#!/usr/bin/env python3
"""Send a simple mail merge from contact-list.xlsx and a Markdown template."""

from __future__ import annotations

import argparse
import csv
import os
import re
import smtplib
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from typing import Iterable

from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv()


REQUIRED_COLUMNS = ("Email",)
PLACEHOLDER_RE = re.compile(r"{{\s*([^{}]+?)\s*}}")
_IF_BLOCK_RE = re.compile(r"\{%\s*if\s+([^%]+?)\s*%\}(.*?)\{%\s*endif\s*%\}", re.DOTALL)
_DIGIT_RE = re.compile(r"\d")
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


@dataclass(frozen=True)
class Contact:
    email: str
    fields: dict[str, str]
    row: int


@dataclass(frozen=True)
class Template:
    subject: str
    cc: tuple[str, ...]
    bcc: tuple[str, ...]
    body: str


@dataclass(frozen=True)
class ValidationIssue:
    row: int
    email: str
    field: str
    message: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Merge Excel columns into a Markdown email template.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""\
Examples:
  python mail_merge.py
  python mail_merge.py --limit 1
  python mail_merge.py --template invite.md
  python mail_merge.py --send --limit 1
  python mail_merge.py --send --delay 5
  python mail_merge.py --template --contacts contacts.csv   (scaffold a new template)

Template front matter:
  ---
  subject: "Hello {{ First Name }} {{ Last Name }}"
  cc: "someone@example.com"
  bcc: "archive@example.com"
  ---

Any Excel column header can be used as a placeholder, for example:
  {{ First Name }}
  {{ Last Name }}
  {{ Email }}
""",
    )
    parser.add_argument(
        "--contacts",
        default="contact-list.xlsx",
        help="Path to the XLSX, ODS, or CSV contact list. Default: contact-list.xlsx",
    )
    parser.add_argument(
        "--template",
        nargs="?",
        default="email-template.md",
        const=None,
        help=(
            "Path to the Markdown email template (default: email-template.md). "
            "Pass the flag with no value to scaffold a new template from --contacts."
        ),
    )
    parser.add_argument(
        "--send",
        action="store_true",
        help="Actually send emails. Omit this for a dry run preview.",
    )
    parser.add_argument(
        "--preview",
        action="store_true",
        help="Preview emails without sending. This is the default behavior.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Optional maximum number of contacts to process.",
    )
    parser.add_argument(
        "--delay",
        type=float,
        default=3.0,
        help="Seconds to wait between sent emails. Default: 3",
    )
    args = parser.parse_args()
    if args.send and args.preview:
        parser.error("--send and --preview cannot be used together")
    return args


def scaffold_template(contacts_path: Path) -> int:
    output_path = Path(contacts_path.stem + "-template.md")
    if output_path.exists():
        print(
            f"Error: {output_path} already exists. Remove it or rename it before scaffolding.",
            file=sys.stderr,
        )
        return 1

    contacts = load_contacts(contacts_path)
    if not contacts:
        print("No contacts found — cannot infer columns.", file=sys.stderr)
        return 1

    columns = list(contacts[0].fields.keys())
    non_email = [c for c in columns if c.lower() != "email"]

    subject_placeholders = " ".join(f"{{{{ {c} }}}}" for c in non_email[:2]) or "Your Subject"
    greeting_field = non_email[0] if non_email else columns[0]
    placeholders_inline = ", ".join(f"{{{{ {c} }}}}" for c in columns)

    content = (
        "---\n"
        f'subject: "{subject_placeholders}"\n'
        'cc: ""\n'
        'bcc: ""\n'
        "---\n"
        f"Hi {{{{ {greeting_field} }}}},\n"
        "\n"
        "Write your message here.\n"
        "\n"
        f"Available placeholders: {placeholders_inline}\n"
        "\n"
        "Regards,\n"
        "\n"
        "[Your Name]\n"
    )

    output_path.write_text(content, encoding="utf-8")
    print(f"Template scaffolded: {output_path}")
    print(f"Columns: {', '.join(columns)}")
    return 0


def load_contacts(path: Path) -> list[Contact]:
    if path.suffix.lower() == ".csv":
        return load_csv_contacts(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return load_excel_contacts(path)
    if path.suffix.lower() == ".ods":
        return load_ods_contacts(path)
    raise ValueError(f"{path} must be an .xlsx, .xlsm, .ods, or .csv file")


def load_excel_contacts(path: Path) -> list[Contact]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration:
        raise ValueError(f"{path} is empty") from None

    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"{path} is missing required column(s): {', '.join(missing)}")

    email_index = headers.index("Email")
    contacts: list[Contact] = []

    for row_number, row in enumerate(rows, start=2):
        fields = {header: cell_text(row, index) for index, header in enumerate(headers)}
        email = cell_text(row, email_index)

        if not any(fields.values()):
            continue
        if not email:
            print(f"Skipping row {row_number}: missing Email", file=sys.stderr)
            continue

        contacts.append(Contact(email=email, fields=fields, row=row_number))

    return contacts


def load_csv_contacts(path: Path) -> list[Contact]:
    with path.open(newline="", encoding="utf-8-sig") as csv_file:
        reader = csv.DictReader(csv_file)
        if not reader.fieldnames:
            raise ValueError(f"{path} is empty")

        headers = [str(value).strip() if value is not None else "" for value in reader.fieldnames]
        missing = [column for column in REQUIRED_COLUMNS if column not in headers]
        if missing:
            raise ValueError(f"{path} is missing required column(s): {', '.join(missing)}")

        contacts: list[Contact] = []
        for row_number, row in enumerate(reader, start=2):
            fields = {
                header: str(row.get(header, "") or "").strip()
                for header in headers
            }
            email = fields["Email"]

            if not any(fields.values()):
                continue
            if not email:
                print(f"Skipping row {row_number}: missing Email", file=sys.stderr)
                continue

            contacts.append(Contact(email=email, fields=fields, row=row_number))

    return contacts


def load_ods_contacts(path: Path) -> list[Contact]:
    try:
        from odf.opendocument import load
        from odf.table import Table, TableCell, TableRow
        from odf.text import P
    except ImportError as exc:
        raise ValueError(
            "Reading .ods files requires odfpy. Run: pip install -r requirements.txt"
        ) from exc

    document = load(str(path))
    tables = document.spreadsheet.getElementsByType(Table)
    if not tables:
        raise ValueError(f"{path} does not contain any sheets")

    rows = ods_table_rows(tables[0], TableRow, TableCell, P)
    if not rows:
        raise ValueError(f"{path} is empty")

    headers = [value.strip() for value in rows[0]]
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"{path} is missing required column(s): {', '.join(missing)}")

    contacts: list[Contact] = []
    for row_number, row in enumerate(rows[1:], start=2):
        fields = {
            header: row[index].strip() if index < len(row) else ""
            for index, header in enumerate(headers)
        }
        email = fields["Email"]

        if not any(fields.values()):
            continue
        if not email:
            print(f"Skipping row {row_number}: missing Email", file=sys.stderr)
            continue

        contacts.append(Contact(email=email, fields=fields, row=row_number))

    return contacts


def ods_table_rows(table: object, table_row_type: object, table_cell_type: object, p_type: object) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.getElementsByType(table_row_type):
        row_repeat = int(row.getAttribute("numberrowsrepeated") or 1)
        values: list[str] = []

        for cell in row.getElementsByType(table_cell_type):
            cell_repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
            text = " ".join(
                node_text(paragraph).strip()
                for paragraph in cell.getElementsByType(p_type)
                if node_text(paragraph).strip()
            )
            values.extend([text] * cell_repeat)

        for _ in range(row_repeat):
            rows.append(values.copy())

    return rows


def node_text(node: object) -> str:
    text_parts: list[str] = []
    for child in getattr(node, "childNodes", []):
        if hasattr(child, "data"):
            text_parts.append(str(child.data))
        else:
            text_parts.append(node_text(child))
    return "".join(text_parts)


def cell_text(row: tuple[object, ...], index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def load_template(path: Path) -> Template:
    text = path.read_text(encoding="utf-8")
    front_matter_values: dict[str, str] = {}
    body = text

    if text.startswith("---\n"):
        _, front_matter, body = text.split("---\n", 2)
        front_matter_values = parse_front_matter(front_matter)

    subject = front_matter_values.get("subject", "")
    if not subject:
        raise ValueError(
            f"{path} must start with front matter containing a subject, for example: "
            'subject: "Hello {{ First Name }}"'
        )

    return Template(
        subject=subject,
        cc=parse_recipients(front_matter_values.get("cc", "")),
        bcc=parse_recipients(front_matter_values.get("bcc", "")),
        body=body.strip() + "\n",
    )


def parse_front_matter(front_matter: str) -> dict[str, str]:
    values: dict[str, str] = {}
    for line in front_matter.splitlines():
        key, sep, value = line.partition(":")
        if sep:
            values[key.strip().lower()] = value.strip().strip("\"'")
    return values


def parse_recipients(value: str) -> tuple[str, ...]:
    if not value:
        return ()
    return tuple(item.strip() for item in value.split(",") if item.strip())


def render(text: str, contact: Contact) -> str:
    def eval_if_block(match: re.Match[str]) -> str:
        field_name = match.group(1).strip()
        if field_name not in contact.fields:
            raise ValueError(
                f"Template uses {{% if {field_name} %}}, but contact list "
                f"does not have a '{field_name}' column"
            )
        return match.group(2) if contact.fields[field_name] else ""

    def replace(match: re.Match[str]) -> str:
        field_name = match.group(1).strip()
        if field_name not in contact.fields:
            raise ValueError(
                f"Template uses {{{{ {field_name} }}}}, but contact-list.xlsx "
                f"does not have a '{field_name}' column"
            )
        return contact.fields[field_name]

    text = _IF_BLOCK_RE.sub(eval_if_block, text)
    rendered = PLACEHOLDER_RE.sub(replace, text)
    return re.sub(r'\n{3,}', '\n\n', rendered)


def preview(
    contact: Contact,
    subject: str,
    cc: tuple[str, ...],
    bcc: tuple[str, ...],
    body: str,
) -> None:
    print("=" * 72)
    print(f"To: {contact.email}")
    if cc:
        print(f"Cc: {', '.join(cc)}")
    if bcc:
        print(f"Bcc: {', '.join(bcc)}")
    print(f"Subject: {subject}")
    print("-" * 72)
    print(body, end="" if body.endswith("\n") else "\n")


def md_to_html(text: str) -> str:
    import markdown
    body_html = markdown.markdown(text, extensions=["nl2br"])
    return f"<html><body>{body_html}</body></html>"


def send_email(
    contact: Contact,
    subject: str,
    cc: tuple[str, ...],
    bcc: tuple[str, ...],
    body: str,
) -> None:
    host = required_env("SMTP_HOST")
    port = int(os.getenv("SMTP_PORT", "587"))
    username = required_env("SMTP_USER")
    password = required_env("SMTP_PASSWORD")
    from_email = os.getenv("FROM_EMAIL", username).strip() or username
    from_name = os.getenv("FROM_NAME", "").strip()
    sender = f"{from_name} <{from_email}>" if from_name else from_email

    message = EmailMessage()
    message["From"] = sender
    message["To"] = contact.email
    if cc:
        message["Cc"] = ", ".join(cc)
    if bcc:
        message["Bcc"] = ", ".join(bcc)
    message["Subject"] = subject
    message.set_content(body)
    message.add_alternative(md_to_html(body), subtype="html")

    with smtplib.SMTP(host, port) as smtp:
        if os.getenv("SMTP_STARTTLS", "true").lower() not in {"0", "false", "no"}:
            smtp.starttls()
        smtp.login(username, password)
        smtp.send_message(message)


def required_env(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise ValueError(f"Missing required environment variable: {name}")
    return value


def limited_contacts(contacts: list[Contact], limit: int | None) -> Iterable[Contact]:
    return contacts[:limit] if limit is not None else contacts


def validate_contacts(contacts: list[Contact]) -> list[ValidationIssue]:
    issues: list[ValidationIssue] = []
    _check_duplicate_emails(contacts, issues)
    _check_email_format(contacts, issues)
    _check_name_fields(contacts, issues)
    return issues


def _check_duplicate_emails(contacts: list[Contact], issues: list[ValidationIssue]) -> None:
    seen: dict[str, int] = {}
    for contact in contacts:
        key = contact.email.lower()
        if key in seen:
            issues.append(ValidationIssue(
                row=contact.row, email=contact.email, field="Email",
                message=f"duplicate of row {seen[key]}",
            ))
        else:
            seen[key] = contact.row


def _check_email_format(contacts: list[Contact], issues: list[ValidationIssue]) -> None:
    for contact in contacts:
        if not _EMAIL_RE.match(contact.email):
            issues.append(ValidationIssue(
                row=contact.row, email=contact.email, field="Email",
                message="invalid format",
            ))


def _check_name_part(row: int, email: str, field: str, value: str, issues: list[ValidationIssue]) -> None:
    if not value:
        issues.append(ValidationIssue(row=row, email=email, field=field, message="empty"))
    elif _DIGIT_RE.search(value):
        issues.append(ValidationIssue(
            row=row, email=email, field=field,
            message=f"contains digits: {value!r}",
        ))


def _check_name_fields(contacts: list[Contact], issues: list[ValidationIssue]) -> None:
    if not contacts:
        return

    cols_lower = {c.lower(): c for c in contacts[0].fields}
    first_col = cols_lower.get("first name")
    last_col = cols_lower.get("last name")
    name_col = cols_lower.get("name")

    if first_col and last_col:
        for contact in contacts:
            first = contact.fields.get(first_col, "").strip()
            last = contact.fields.get(last_col, "").strip()
            _check_name_part(contact.row, contact.email, first_col, first, issues)
            _check_name_part(contact.row, contact.email, last_col, last, issues)
            if first and last and first.lower() == last.lower():
                issues.append(ValidationIssue(
                    row=contact.row, email=contact.email,
                    field=f"{first_col}/{last_col}",
                    message=f"first and last name are identical: {first!r}",
                ))
    elif name_col:
        for contact in contacts:
            name = contact.fields.get(name_col, "").strip()
            if not name:
                issues.append(ValidationIssue(row=contact.row, email=contact.email, field=name_col, message="empty"))
                continue
            if _DIGIT_RE.search(name):
                issues.append(ValidationIssue(
                    row=contact.row, email=contact.email, field=name_col,
                    message=f"contains digits: {name!r}",
                ))
            parts = name.split()
            if len(parts) == 1:
                issues.append(ValidationIssue(
                    row=contact.row, email=contact.email, field=name_col,
                    message=f"only one word — missing first or last name: {name!r}",
                ))
            elif parts[0].lower() == parts[-1].lower():
                issues.append(ValidationIssue(
                    row=contact.row, email=contact.email, field=name_col,
                    message=f"first and last word are identical: {name!r}",
                ))


def prompt_continue(issues: list[ValidationIssue]) -> bool:
    """Print all validation issues and ask the user to exit or continue."""
    line = "─" * 72
    print(f"\n{line}")
    print(f"  {len(issues)} data quality issue(s) found\n")
    field_w = max(len(iss.field) for iss in issues)
    for iss in issues:
        print(f"  row {iss.row:>3}  {iss.email:<40}  {iss.field:<{field_w}}  {iss.message}")
    print(f"{line}\n")
    print("  1. Exit    — fix the contact file and re-run")
    print("  2. Continue — proceed anyway (duplicate/invalid emails will be skipped)\n")

    if not sys.stdin.isatty():
        print("Non-interactive mode — exiting by default.", file=sys.stderr)
        return False

    while True:
        try:
            choice = input("Choose [1/2]: ").strip()
        except (EOFError, KeyboardInterrupt):
            print()
            return False
        if choice == "1":
            return False
        if choice == "2":
            return True
        print("Please enter 1 or 2.")


def clean_contacts(contacts: list[Contact]) -> list[Contact]:
    """Remove contacts with duplicate or malformed email addresses (keep first occurrence)."""
    seen: set[str] = set()
    result: list[Contact] = []
    for contact in contacts:
        if not _EMAIL_RE.match(contact.email):
            print(f"Skipping {contact.email!r}: invalid email format.", file=sys.stderr)
            continue
        key = contact.email.lower()
        if key in seen:
            print(f"Skipping {contact.email!r}: duplicate.", file=sys.stderr)
            continue
        seen.add(key)
        result.append(contact)
    return result


def write_audit_trail(
    path: Path,
    sent_contacts: list[Contact],
    template_name: str,
    run_date: str,
    run_time: str,
) -> None:
    if not sent_contacts:
        return
    sent_col = f"Sent - {run_date}"
    tmpl_col = f"Template - {run_date}"
    ext = path.suffix.lower()
    try:
        if ext in {".xlsx", ".xlsm"}:
            _write_audit_xlsx(path, sent_contacts, sent_col, tmpl_col, run_time, template_name)
        elif ext == ".csv":
            _write_audit_csv(path, sent_contacts, sent_col, tmpl_col, run_time, template_name)
        else:
            print(f"Note: audit trail write-back is not supported for {ext} files.", file=sys.stderr)
            return
        print(f"Audit trail written to {path}.")
    except Exception as exc:
        print(f"Warning: could not write audit trail to {path}: {exc}", file=sys.stderr)


def _write_audit_xlsx(
    path: Path,
    sent_contacts: list[Contact],
    sent_col: str,
    tmpl_col: str,
    run_time: str,
    template_name: str,
) -> None:
    wb = load_workbook(path)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    def col_index(name: str) -> int:
        if name in headers:
            return headers.index(name) + 1
        headers.append(name)
        ws.cell(row=1, column=len(headers), value=name)
        return len(headers)

    sent_idx = col_index(sent_col)
    tmpl_idx = col_index(tmpl_col)

    for contact in sent_contacts:
        ws.cell(row=contact.row, column=sent_idx, value=run_time)
        ws.cell(row=contact.row, column=tmpl_idx, value=template_name)

    wb.save(path)


def _write_audit_csv(
    path: Path,
    sent_contacts: list[Contact],
    sent_col: str,
    tmpl_col: str,
    run_time: str,
    template_name: str,
) -> None:
    with path.open(newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))

    if not rows:
        return

    headers = rows[0]

    def col_index(name: str) -> int:
        if name in headers:
            return headers.index(name)
        headers.append(name)
        for row in rows[1:]:
            row.append("")
        return len(headers) - 1

    sent_idx = col_index(sent_col)
    tmpl_idx = col_index(tmpl_col)

    sent_row_nums = {contact.row for contact in sent_contacts}
    for file_row, row in enumerate(rows[1:], start=2):
        if file_row in sent_row_nums:
            while len(row) <= max(sent_idx, tmpl_idx):
                row.append("")
            row[sent_idx] = run_time
            row[tmpl_idx] = template_name

    with path.open("w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(rows)


def main() -> int:
    args = parse_args()

    if args.template is None:
        return scaffold_template(Path(args.contacts))

    contacts = load_contacts(Path(args.contacts))
    template = load_template(Path(args.template))

    if not contacts:
        print("No contacts found.")
        return 0

    issues = validate_contacts(contacts)
    if issues:
        if not prompt_continue(issues):
            return 1
        contacts = clean_contacts(contacts)
        if not contacts:
            print("No valid contacts remain after cleanup.")
            return 1
        print(f"Proceeding with {len(contacts)} contact(s).\n")

    now = datetime.now()
    run_date = now.strftime("%Y-%m-%d")
    run_time = now.strftime("%H:%M")

    processed = 0
    sent_contacts: list[Contact] = []
    selected_contacts = list(limited_contacts(contacts, args.limit))

    for index, contact in enumerate(selected_contacts, start=1):
        subject = render(template.subject, contact)
        cc = tuple(render(value, contact) for value in template.cc)
        bcc = tuple(render(value, contact) for value in template.bcc)
        body = render(template.body, contact)

        if args.send:
            send_email(contact, subject, cc, bcc, body)
            print(f"Sent: {contact.email}")
            sent_contacts.append(contact)
            if args.delay > 0 and index < len(selected_contacts):
                print(f"Waiting {args.delay:g}s before next email...")
                time.sleep(args.delay)
        else:
            preview(contact, subject, cc, bcc, body)

        processed += 1

    mode = "sent" if args.send else "previewed"
    print(f"\nDone: {mode} {processed} email(s).")
    if not args.send:
        print("Dry run only. Add --send after Gmail SMTP is configured.")
    else:
        write_audit_trail(
            Path(args.contacts), sent_contacts,
            Path(args.template).name, run_date, run_time,
        )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        raise SystemExit(1)
